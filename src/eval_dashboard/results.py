"""Persist one evaluation run to a document (and JSON) on disk.

The browser POSTs the obstacle timings plus the 5-question survey; this turns
that dict into a human-readable report. It writes a real Word `.docx` when
`python-docx` is installed, and otherwise falls back to a plain-text `.txt`
with the same content -- so a run is never lost just because an optional
dependency is missing. A `.json` sibling is always written for later analysis.

`save_results` returns the path of the primary document it wrote.
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

# The 5 survey prompts, in order. Kept here so the report and the front-end
# agree on wording; the browser sends answers keyed q1..q5.
SURVEY_QUESTIONS = [
    "How well were you able to click?",
    "How well was your hand tracked?",
    "How easy was it to use overall?",
    "How easy were the controls to understand?",
    "How well did it do what you wanted?",
]


def _fmt_seconds(ms: float | None) -> str:
    if ms is None:
        return "n/a"
    return f"{ms / 1000:.2f} s"


def _report_lines(data: dict) -> list[tuple[str, str]]:
    """Flatten a run into (style, text) lines. style in {h1,h2,p,kv}."""
    ob1 = data.get("ob1", {}) or {}
    ob2 = data.get("ob2", {}) or {}
    ob3 = data.get("ob3", {}) or {}
    survey = data.get("survey", {}) or {}
    lines: list[tuple[str, str]] = [("h1", "Hand-Tracking Evaluation Report")]

    lines.append(("p", f"Run started: {data.get('startedAt', 'n/a')}"))
    lines.append(("p", f"Run finished: {data.get('finishedAt', 'n/a')}"))
    scr = data.get("screen") or {}
    if scr:
        lines.append(("p", f"Viewport: {scr.get('w', '?')} x {scr.get('h', '?')} px"))

    lines.append(("h2", "Obstacle 1 - Precision (buttons 1-5)"))
    lines.append(("kv", f"Time (click 1 -> click 5): {_fmt_seconds(ob1.get('timeMs'))}"))
    lines.append(("kv", f"Total incl. acquiring button 1: {_fmt_seconds(ob1.get('totalMs'))}"))
    if ob1.get("splitsMs"):
        splits = ", ".join(f"{s / 1000:.2f}s" for s in ob1["splitsMs"])
        lines.append(("kv", f"Per-button splits: {splits}"))

    lines.append(("h2", "Obstacle 2 - Rapid clicks (down x7)"))
    lines.append(("kv", f"Time (7 clicks): {_fmt_seconds(ob2.get('timeMs'))}"))
    lines.append(("kv", f"Clicks: {ob2.get('clicks', 'n/a')}"))
    if ob2.get("timeMs") and ob2.get("clicks"):
        rate = ob2["clicks"] / (ob2["timeMs"] / 1000)
        lines.append(("kv", f"Click rate: {rate:.2f} clicks/s"))

    lines.append(("h2", "Obstacle 3 - Line tracer"))
    lines.append(("kv", f"Time: {_fmt_seconds(ob3.get('timeMs'))}"))
    acc = ob3.get("accuracy")
    lines.append(("kv", f"Accuracy score: {acc:.1f} / 100" if acc is not None else "Accuracy score: n/a"))
    if ob3.get("meanDevPx") is not None:
        lines.append(("kv", f"Mean deviation from line: {ob3['meanDevPx']:.1f} px"))
    if ob3.get("coverage") is not None:
        lines.append(("kv", f"Path coverage: {ob3['coverage'] * 100:.0f}%"))

    lines.append(("h2", "Survey (1 = poor, 5 = excellent)"))
    total, count = 0, 0
    for i, q in enumerate(SURVEY_QUESTIONS, start=1):
        ans = survey.get(f"q{i}")
        lines.append(("kv", f"{i}. {q}  ->  {ans if ans is not None else 'n/a'}"))
        if isinstance(ans, (int, float)):
            total += ans
            count += 1
    if count:
        lines.append(("kv", f"Average rating: {total / count:.2f} / 5"))
    return lines


def _write_docx(path: Path, lines: list[tuple[str, str]]) -> bool:
    try:
        from docx import Document
    except ImportError:
        return False
    doc = Document()
    for style, text in lines:
        if style == "h1":
            doc.add_heading(text, level=0)
        elif style == "h2":
            doc.add_heading(text, level=1)
        elif style == "kv":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)
    doc.save(str(path))
    return True


def _write_txt(path: Path, lines: list[tuple[str, str]]) -> None:
    out: list[str] = []
    for style, text in lines:
        if style == "h1":
            out += [text, "=" * len(text), ""]
        elif style == "h2":
            out += ["", text, "-" * len(text)]
        elif style == "kv":
            out.append(f"  - {text}")
        else:
            out.append(text)
    path.write_text("\n".join(out) + "\n", encoding="utf-8")


def summarize_run(json_path: str | Path) -> dict | None:
    """Compact summary of one saved run for the Overview tab (None if unreadable).

    Pulls the headline number from each obstacle + the survey average, and the
    model/strategy `meta` block if the run recorded one (older runs won't have it).
    """
    json_path = Path(json_path)
    try:
        data = json.loads(json_path.read_text(encoding="utf-8"))
    except Exception:
        return None
    ob1 = data.get("ob1") or {}
    ob2 = data.get("ob2") or {}
    ob3 = data.get("ob3") or {}
    survey = data.get("survey") or {}
    meta = data.get("meta") or {}
    svals = [v for v in survey.values() if isinstance(v, (int, float))]
    survey_avg = round(sum(svals) / len(svals), 2) if svals else None
    return {
        "id": json_path.stem,
        "startedAt": data.get("startedAt"),
        "finishedAt": data.get("finishedAt"),
        "ob1Ms": ob1.get("timeMs"),
        "ob2Ms": ob2.get("timeMs"),
        "ob2Clicks": ob2.get("clicks"),
        "ob3Accuracy": ob3.get("accuracy"),
        "ob3Coverage": ob3.get("coverage"),
        "ob3MeanDevPx": ob3.get("meanDevPx"),
        "surveyAvg": survey_avg,
        "model": meta.get("model"),
        "strategy": meta.get("strategy"),
        "device": meta.get("device"),
    }


def save_results(data: dict, out_dir: str | Path = "eval_results") -> Path:
    """Write the run to a doc file (+ JSON). Returns the document path.

    Tries `.docx` (python-docx); falls back to `.txt`. The JSON sibling always
    holds the raw payload for later analysis.
    """
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = out_dir / f"eval_{stamp}"

    # Raw payload for downstream analysis (always written).
    base.with_suffix(".json").write_text(
        json.dumps(data, indent=2, default=str), encoding="utf-8"
    )

    lines = _report_lines(data)
    docx_path = base.with_suffix(".docx")
    if _write_docx(docx_path, lines):
        return docx_path
    txt_path = base.with_suffix(".txt")
    _write_txt(txt_path, lines)
    return txt_path
